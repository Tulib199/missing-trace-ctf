from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x4E, 0x79)
INK = RGBColor(0x23, 0x23, 0x23)
GRAY = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xA6, 0x1B, 0x1B)

OUT_DIR = r"c:\Users\Tulib\Desktop\Cursor\Residency IWU - Capture the flag"


def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Emu(640080)
    sec.top_margin = sec.bottom_margin = Emu(548640)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.paragraph_format.space_after = Pt(4)
    return doc


def shade(element_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element_pr.append(shd)


def title_block(doc, title, subtitle, lines):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.font.name, r.font.size, r.bold, r.font.color.rgb = "Aptos Display", Pt(20), True, NAVY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.font.size, r.bold, r.font.color.rgb = Pt(12), True, INK
    for line in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.add_run(line).font.size = Pt(9)
    doc.add_paragraph()


def heading(doc, text, size=13):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name, r.font.size, r.bold, r.font.color.rgb = "Aptos Display", Pt(size), True, NAVY
    return p


def para(doc, text, bold_lead=None, italic=False, color=INK, size=10):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(size), color
    r = p.add_run(text)
    r.italic, r.font.size, r.font.color.rgb = italic, Pt(size), color
    return p


def bullet(doc, text, bold_lead=None, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(2)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold, r.font.size = True, Pt(10)
    p.add_run(text).font.size = Pt(10)
    return p


def code_block(doc, lines):
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(4) if i == 0 else Pt(0)
        pf.space_after = Pt(4) if i == len(lines) - 1 else Pt(0)
        pf.left_indent = Inches(0.1)
        shade(p._p.get_or_add_pPr(), "F2F2F2")
        r = p.add_run(line)
        r.font.name, r.font.size = "Consolas", Pt(8)
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")


def table(doc, headers, rows, widths=None, font=8.5, header_fill="1F4E79"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        shade(cell._tc.get_or_add_tcPr(), header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(font), RGBColor(0xFF, 0xFF, 0xFF)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            if ri % 2 == 1:
                shade(cells[i]._tc.get_or_add_tcPr(), "EAF1F8")
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.add_run(str(val)).font.size = Pt(font)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    if len(rows) <= 12:
        for row in t.rows[:-1]:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def answer_lines(doc, label, n=2):
    para(doc, label, bold_lead=None)
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.add_run("_" * 98).font.color.rgb = GRAY


# ---------------------------------------------------------------- packet

def build_packet():
    doc = new_doc()
    title_block(
        doc,
        "The Missing Trace",
        "A Real-Time Data Integration and Process Optimization CTF Challenge",
        [
            "CITM-501: Information Systems Infrastructure",
            "Assignment 6.6 \u2013 Capture the Flag Activity",
            "Challenge Author: Meghana Sinha \u2013 2953222",
            "Team size: 3\u20134 players  |  Target completion time: 45 minutes  |  Version 2",
        ],
    )

    heading(doc, "Challenge Objective")
    para(doc, "NorthBridge Health Supply is a fictional distributor of serialized medical products. "
              "Every pallet received at the warehouse is scanned, posted to the ERP system, and passed in "
              "real time through an integration layer to a cloud traceability platform. Regulators and "
              "hospital customers rely on that platform to prove where each serialized item has been.")
    para(doc, "At 09:30 the following ticket arrives at the service desk:")
    code_block(doc, [
        "INC-20931  Priority: High  Opened: 09:30  Caller: Warehouse Operations Lead",
        "\"The traceability dashboard is showing 8 receipts missing since about 09:10. They are all in",
        " the ERP. My team has been resending and rescanning to try to push them through. Please fix",
        " this before the 11:00 hospital shipment cut-off.\"",
    ])
    para(doc, "Your team is the integration support group. Using only the evidence in this packet, determine "
              "which component failed, exactly when it started failing, how many unique business "
              "transactions actually need to be recovered, and which long-term design prevents a repeat.")

    heading(doc, "How to Play as a Team")
    para(doc, "The evidence is split into four work streams. Assign one role per person (combine roles 3 and 4 "
              "for a team of three). Each role contributes one part of the flag, so no one can finish alone. "
              "Plan to work in parallel for the first 20 minutes, then regroup and compare findings.")
    table(doc, ["Role", "Primary Evidence", "Delivers to the Team"], [
        ["1. Data Reconciliation Analyst", "A (Transactions), G (Scanner & Ops Log)",
         "Which missing receipts are unique, which are duplicates, and which never reached the integration layer"],
        ["2. Log Analyst", "B (Hub Log), C (Proxy Log), H (Change Calendar)",
         "The exact time the root-cause component began failing, and what changed"],
        ["3. Infrastructure Analyst", "D (Connectivity Tests), E (Vendor Status)",
         "Which network path and component is broken, and which are healthy"],
        ["4. Solutions Architect", "F (Architecture Options) + findings from roles 1\u20133",
         "The option that prevents recurrence without creating new data-quality risk"],
    ], widths=[1.7, 2.2, 3.3])

    heading(doc, "Rules")
    for txt in [
        "Use only the evidence included in this packet. No outside lookups are needed.",
        "The component that reports an error is not necessarily the component that caused it. Compare evidence across layers.",
        "A \u201cmissing\u201d receipt on the dashboard is not automatically a failed integration message. Trace each one.",
        "Separate retries, manual resends, and rescans from unique business transactions.",
        "Your final answer is a single flag in the exact format shown at the end. It is checked character-for-character.",
        "Two optional hints are available at the end of the packet. Note the time if you use them.",
    ]:
        bullet(doc, txt)

    heading(doc, "Current System Flow")
    table(doc, ["1", "2", "3", "4", "5", "6"], [
        ["Warehouse Scanner", "ERP Receiving", "Integration Hub", "Legacy App Proxy", "Traceability Cloud (TraceCloud)", "Operations Dashboard"],
        ["Wi-Fi VLAN 30 \u2192 HTTPS", "Posts receipt, creates a Match ID", "Holds and retries messages (3\u00d7, 30 s)", "Forwards hub \u2192 cloud, port 8443", "Vendor API, port 443", "Reads TraceCloud data"],
    ], widths=[1.2] * 6, font=8)
    para(doc, "ERP creates a Match ID from the purchase order and the serial numbers on the pallet. "
              "If two rows have the same Match ID, they are the same physical pallet \u2014 even if the transaction IDs are different.", size=9, color=GRAY)

    heading(doc, "Business Impact")
    para(doc, "Operations uses the traceability dashboard to confirm that serialized products were received and are "
              "available for downstream picking and shipping. Products that are not visible cannot ship. When receipts "
              "appear missing, staff tend to resend or rescan, which creates duplicate serial records in the cloud "
              "platform if those duplicates are later replayed. Duplicate serial records must be manually investigated "
              "and can trigger compliance holds.")

    heading(doc, "Evidence A \u2013 Transaction Timeline (ERP + Dashboard export, 09:29)")
    table(doc, ["Txn ID", "ERP Time", "ERP Status", "PO", "Match ID", "Hub Status", "Cloud Ack", "Dashboard"], [
        ["T1001", "09:02:14", "Posted", "PO-5510", "RX-88410", "Forwarded", "ACK-200", "Visible"],
        ["T1002", "09:05:29", "Posted", "PO-5511", "RX-88411", "Forwarded", "ACK-200", "Visible"],
        ["T1003", "09:06:11", "Posted", "PO-5512", "RX-88412", "Parked", "None", "Missing"],
        ["T1004", "09:08:08", "Posted", "PO-5513", "RX-88415", "Parked", "None", "Missing"],
        ["T1005", "09:10:03", "Rejected (VAL-112)", "PO-5514", "\u2014", "Not received", "None", "Missing"],
        ["T1006", "09:11:30", "Posted", "PO-5515", "RX-88419", "Parked", "None", "Missing"],
        ["T1007", "09:14:17", "Posted", "PO-5516", "RX-88421", "Parked", "None", "Missing"],
        ["T1008", "09:18:04", "Posted", "PO-5515", "RX-88419", "Parked", "None", "Missing"],
        ["T1009", "09:21:39", "Posted", "PO-5517", "RX-88426", "Parked", "None", "Missing"],
        ["T1010", "09:23:01", "Posted", "PO-5512", "RX-88412", "Parked", "None", "Missing"],
        ["TEST-DIRECT", "09:25:10", "N/A", "\u2014", "\u2014", "Health test", "HTTP 200", "N/A"],
    ], widths=[0.85, 0.75, 1.15, 0.75, 0.85, 0.95, 0.8, 0.8], font=8)
    para(doc, "VAL-112 = the ERP rejected the receipt because the lot number was blank. Rejected receipts are never sent to the hub. "
              "Parked means the hub tried three times, failed, and is holding the message.",
         size=9, color=GRAY)

    heading(doc, "Evidence B \u2013 Integration Hub Log")
    code_block(doc, [
        "08:00:00  CONFIG  retry_max=3  retry_interval=30s  dedupe_check=OFF  parked_queue_alert=OFF",
        "09:02:15  MSG T1001  key=RX-88410  route=legacy-proxy:8443  attempt=1  result=HTTP_200  ack=ACK-200",
        "09:05:30  MSG T1002  key=RX-88411  route=legacy-proxy:8443  attempt=1  result=HTTP_200  ack=ACK-200",
        "09:06:12  MSG T1003  key=RX-88412  route=legacy-proxy:8443  attempt=1  result=HTTP_504_GATEWAY_TIMEOUT",
        "09:06:42  MSG T1003  key=RX-88412  route=legacy-proxy:8443  attempt=2  result=HTTP_504_GATEWAY_TIMEOUT",
        "09:07:12  MSG T1003  key=RX-88412  route=legacy-proxy:8443  attempt=3  result=RETRY_EXHAUSTED  status=PARKED",
        "09:08:09  MSG T1004  key=RX-88415  route=legacy-proxy:8443  attempts=3  result=RETRY_EXHAUSTED  status=PARKED",
        "09:11:31  MSG T1006  key=RX-88419  route=legacy-proxy:8443  attempts=3  result=RETRY_EXHAUSTED  status=PARKED",
        "09:14:18  MSG T1007  key=RX-88421  route=legacy-proxy:8443  attempts=3  result=RETRY_EXHAUSTED  status=PARKED",
        "09:18:05  MSG T1008  key=RX-88419  source=MANUAL_RESEND  attempts=3  result=RETRY_EXHAUSTED  status=PARKED",
        "09:21:40  MSG T1009  key=RX-88426  route=legacy-proxy:8443  attempts=3  result=RETRY_EXHAUSTED  status=PARKED",
        "09:23:02  MSG T1010  key=RX-88412  source=ERP_RECEIPT  attempts=3  result=RETRY_EXHAUSTED  status=PARKED",
        "09:25:10  TEST-DIRECT  route=tracecloud-api:443 (temporary rule FW-TMP-77)  result=HTTP_200  health=OK",
    ])

    heading(doc, "Evidence C \u2013 Legacy App Proxy Log")
    code_block(doc, [
        "08:52:10  CHG-4471  OS security patch applied; relay configuration reloaded",
        "08:52:11  truststore=/opt/lap/certs rebuilt from baseline image  entries=112",
        "08:55:47  graceful restart completed; existing upstream pool connections retained  max_age=600s",
        "09:02:15  relay T1001  upstream=tracecloud-api:443  pooled_conn=7  status=200",
        "09:05:30  relay T1002  upstream=tracecloud-api:443  pooled_conn=3  status=200",
        "09:05:47  pool=tracecloud  all pooled connections reached max_age; opening new upstream connections",
        "09:05:47  upstream TLS handshake failed: unable to get local issuer certificate (TraceCloud Issuing CA G2)",
        "09:05:47  connector_pool=tracecloud  status=UNAVAILABLE",
        "09:06:12  inbound session accepted from integration-hub",
        "09:06:42  upstream connection not established before timeout (30s); returned 504 to client",
        "09:08:09  upstream connection not established before timeout (30s); returned 504 to client",
        "09:11:31  upstream connection not established before timeout (30s); returned 504 to client",
        "   ...    (same pattern repeats for every relay request through 09:23:02)",
    ])

    heading(doc, "Evidence D \u2013 Connectivity Tests (Network Team, 09:25\u201309:28)")
    table(doc, ["Source", "Destination", "Test", "Result"], [
        ["Warehouse scanner (VLAN 30)", "ERP Receiving :443", "HTTPS", "Success"],
        ["ERP server", "Integration Hub :443", "TCP", "Success"],
        ["Integration Hub", "Legacy App Proxy :8443", "TCP", "Success (4 ms)"],
        ["Integration Hub", "Legacy App Proxy :8443", "POST /relay/tracecloud", "HTTP 504 after 30 s"],
        ["Legacy App Proxy", "TraceCloud API :443", "TCP", "Success (22 ms)"],
        ["Legacy App Proxy", "TraceCloud API :443", "TLS handshake", "Failed \u2013 certificate verify failed (unknown issuer)"],
        ["Integration Hub", "TraceCloud API :443", "HTTPS health check via FW-TMP-77", "HTTP 200 \u2013 certificate chain valid (Issuing CA G2)"],
        ["User workstation", "Operations Dashboard", "HTTPS", "Success"],
    ], widths=[1.6, 1.6, 1.8, 2.2])

    heading(doc, "Evidence E \u2013 External Platform Status (TraceCloud Status Page, 09:27)")
    para(doc, "All systems operational. No active incident. API health endpoint responding normally. No maintenance "
              "window in progress. Inbound processing normal for all regions during the last 24 hours.")
    para(doc, "Notice (posted Sep 28): TraceCloud replaced its security certificate with a newer one called \u201cTraceCloud "
              "Issuing CA G2.\u201d Any customer that keeps its own list of trusted certificates must add the new G2 "
              "certificate. Everyday browsers and public lists already include it.")

    heading(doc, "Evidence F \u2013 Architecture Options")
    for lead, txt in [
        ("Option A \u2013 ", "Keep sending data through the old proxy. Give the hub more time (30 seconds to 2 minutes) and "
                          "more retries (3 to 10) so short delays can clear on their own."),
        ("Option B \u2013 ", "Connect the ERP straight to the cloud and shut down both the hub and the old proxy, so there "
                          "are fewer systems that can break."),
        ("Option C \u2013 ", "Keep the hub so IT can still hold and retry messages. Stop using the old proxy and send data "
                          "from the hub straight to the cloud on an approved secure connection. Ignore repeat sendings of "
                          "the same pallet so resends do not create duplicate records, and alert the team when messages "
                          "sit unsent for too long."),
        ("Option D \u2013 ", "Keep the hub and stop using the old proxy; send data from the hub straight to the cloud. Add a "
                          "job that automatically resends every stuck message every 15 minutes."),
    ]:
        bullet(doc, txt, bold_lead=lead)

    heading(doc, "Evidence G \u2013 Warehouse Scanner & Operations Log")
    code_block(doc, [
        "09:06:05  WH-SCN-02  user=wh.recv1  PO-5512  pallet scanned  serials=48  lot=L24-0912",
        "09:10:01  WH-SCN-04  user=wh.recv3  PO-5514  pallet scanned  serials=36  lot=<blank>  label damaged",
        "09:11:24  WH-SCN-02  user=wh.recv1  PO-5515  pallet scanned  serials=60  lot=L24-0915",
        "09:17:40  OPS-DESK   user=wh.ops2   \"PO-5515 not on dashboard - used ERP Resend button\"",
        "09:20:55  WH-SCN-04  user=wh.recv3  PO-5512  pallet RESCANNED  serials=48  lot=L24-0912",
        "                     note=\"not visible on dashboard, scanning again\"",
        "09:24:30  OPS-DESK   user=wh.ops2   \"Told team to stop rescanning until IT responds\"",
    ])

    heading(doc, "Evidence H \u2013 Change Calendar (Today)")
    table(doc, ["Change", "Window", "System", "Description", "Status"], [
        ["CHG-4460", "08:30\u201308:40", "Firewall", "Rule update on warehouse wireless VLAN 30 (scanner network)", "Completed"],
        ["CHG-4471", "08:50\u201309:00", "Legacy App Proxy", "Monthly OS security patch and config reload. Risk: Low. \u201cNo user impact expected.\u201d", "Completed"],
        ["CHG-4475", "22:00\u201323:00", "ERP", "Monthly database index maintenance", "Scheduled"],
        ["INC-20410", "Sep 28", "Legacy App Proxy", "Manually imported TraceCloud Issuing CA G2 into proxy truststore after vendor rotation", "Closed"],
    ], widths=[0.8, 0.95, 1.15, 3.45, 0.85])

    heading(doc, "Your Tasks")
    for lead, txt in [
        ("Root cause. ", "Identify the component that is the root cause of the failure (use the component codes below)."),
        ("Onset. ", "Determine the time (HH:MM, 24-hour) at which the root-cause component itself first became unable to deliver data."),
        ("Recovery scope. ", "Determine how many unique business transactions must be replayed through the integration "
                             "layer. Do not count resends, rescans, or receipts that never reached the hub."),
        ("Immediate recovery. ", "Describe the safest immediate recovery action, including what to do with the duplicate "
                                 "and non-integration items."),
        ("Long-term design. ", "Choose the architecture option (A, B, C, or D) that best prevents recurrence while "
                               "preserving real-time processing and data integrity."),
        ("Control. ", "State one monitoring or control improvement that would have detected this problem before operations did."),
    ]:
        bullet(doc, txt, bold_lead=lead, style="List Number")

    heading(doc, "Component Codes")
    table(doc, ["Component", "Code"], [
        ["Warehouse Scanner", "SCANNER"], ["Firewall", "FIREWALL"], ["ERP Receiving", "ERP"],
        ["Integration Hub", "HUB"], ["Legacy App Proxy", "LEGACYPROXY"],
        ["Traceability Cloud", "TRACECLOUD"], ["Operations Dashboard", "DASHBOARD"],
    ], widths=[2.5, 1.8], font=9)

    heading(doc, "Final Flag Format")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FLAG{<COMPONENT_CODE>_<ONSET_HHMM>_<UNIQUE_REPLAY_COUNT>_<OPTION_LETTER>}")
    r.font.name, r.font.size, r.bold, r.font.color.rgb = "Consolas", Pt(11), True, NAVY
    para(doc, "Formatting rules: uppercase, no spaces, no colon in the time. Example of the format only (not the "
              "answer): FLAG{ERP_1342_3_A}.", size=9.5)

    heading(doc, "Team Answer Sheet")
    para(doc, "Submit the flag for verification. Complete the reasoning below so the playtest reviewer can confirm "
              "the flag came from analysis rather than guessing.", size=9.5)
    answer_lines(doc, "Team members and roles:", 1)
    answer_lines(doc, "1. Root-cause component and the two strongest pieces of evidence:", 2)
    answer_lines(doc, "2. Onset time and the log line that proves it:", 1)
    answer_lines(doc, "3. Transactions to replay (list the IDs) and transactions excluded (with reason):", 2)
    answer_lines(doc, "4. Immediate recovery action:", 2)
    answer_lines(doc, "5. Architecture option and why the other three are weaker:", 2)
    answer_lines(doc, "6. Monitoring / control improvement:", 1)
    answer_lines(doc, "FINAL FLAG:", 1)
    answer_lines(doc, "Start time / finish time / hints used:", 1)

    heading(doc, "Optional Hints")
    para(doc, "Hint 1 (after 15 minutes): ", bold_lead=None, size=10)
    doc.paragraphs[-1].runs[0].bold = True
    doc.paragraphs[-1].add_run("Line up the Match ID column. Two different transaction IDs can describe the same "
                               "physical pallet. Also check whether every missing receipt actually reached the hub.").font.size = Pt(10)
    para(doc, "Hint 2 (after 30 minutes): ", size=10)
    doc.paragraphs[-1].runs[0].bold = True
    doc.paragraphs[-1].add_run("The hub records when it noticed the failure, not when the failure began. Find the "
                               "earliest line where the failing component reports its own problem, and ask why it "
                               "happened then rather than right after the restart.").font.size = Pt(10)

    heading(doc, "References")
    para(doc, "CTF101. (n.d.). CTF 101. https://ctf101.org/", size=9.5)
    para(doc, "CTFd. (n.d.). What is Capture the Flag? https://ctfd.io/whats-a-ctf/", size=9.5)

    path = OUT_DIR + r"\Assignment_6.6_CTF_Challenge_Packet_MS_v2.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------- author key

def build_author():
    doc = new_doc()
    title_block(
        doc,
        "Assignment 6.6 \u2013 Author Materials",
        "Answer Key, Playtest Form, Revision Log, and Debrief",
        ["CITM-501: Information Systems Infrastructure", "Meghana Sinha \u2013 2953222",
         "Challenge: The Missing Trace (Version 2)"],
    )
    p = para(doc, "PRIVATE \u2013 Do not provide Sections A and B to the playtest team before they submit a flag.",
             color=RED, size=10)
    p.runs[0].bold = True

    heading(doc, "A. Author Answer Key")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Correct flag:  FLAG{LEGACYPROXY_0905_5_C}")
    r.font.name, r.font.size, r.bold, r.font.color.rgb = "Consolas", Pt(13), True, NAVY

    bullet(doc, "Legacy App Proxy (LEGACYPROXY). CHG-4471 rebuilt the proxy truststore from a baseline image, which "
                "removed the TraceCloud Issuing CA G2 intermediate that had been imported manually under INC-20410. "
                "The proxy can reach TraceCloud over TCP but cannot complete the TLS handshake. The hub reaches the same "
                "API directly with a valid chain, and the vendor reports no incident.", bold_lead="Root cause: ")
    bullet(doc, "09:05 (09:05:47). The graceful restart at 08:55:47 kept existing pooled connections alive for "
                "max_age=600 s. At 09:05:47 those connections expired, new TLS handshakes failed, and the proxy marked "
                "the connector UNAVAILABLE. T1002 at 09:05:30 still succeeded on a pooled connection. Common traps: 0906 "
                "(the hub's first 504), 0855 (restart), 0852 (patch), and 0910 (the ticket's estimate).",
           bold_lead="Onset: ")
    bullet(doc, "5 unique transactions: T1003, T1004, T1006, T1007, T1009.", bold_lead="Replay count: ")
    table(doc, ["Excluded", "Reason"], [
        ["T1001, T1002", "Succeeded (ACK-200, visible)"],
        ["T1005", "Rejected by ERP validation (VAL-112, blank lot number). Never reached the hub. Must be corrected at the source, not replayed."],
        ["T1008", "Manual resend of T1006 (same dedupe key RX-88419, source=MANUAL_RESEND)"],
        ["T1010", "Warehouse rescan of the T1003 pallet (same dedupe key RX-88412, Evidence G at 09:20:55)"],
        ["TEST-DIRECT", "Health check only, not a business transaction"],
    ], widths=[1.3, 5.9], font=9)
    bullet(doc, "Restore trust on the delivery path. Either re-import the G2 intermediate into the proxy truststore "
                "(or roll back the CHG-4471 truststore) and reset the connector pool, or, under an emergency change, "
                "send traffic over the already-tested direct hub-to-cloud route (FW-TMP-77). Then replay only the five "
                "unique messages once, cancel the parked duplicates T1008 and T1010 (and void duplicate ERP receipt T1010), "
                "send T1005 back to receiving to correct the lot number, confirm ACKs and dashboard visibility, and tell "
                "the warehouse to stop resending and rescanning.", bold_lead="Immediate recovery (any equivalent answer is accepted): ")
    bullet(doc, "Option C.", bold_lead="Long-term design: ")
    table(doc, ["Option", "Why it is weaker / correct"], [
        ["A", "Waits longer on a failure that will never clear. A TLS trust failure is not a timeout problem, and "
              "longer retries delay detection."],
        ["B", "Removes the hub's queueing, retry, transformation, and monitoring. ERP becomes tightly coupled to the "
              "vendor API with no controlled place to park and replay failures."],
        ["C", "Correct. Removes the unnecessary proxy hop, keeps the hub so IT can still hold and retry messages, "
              "ignores repeat sendings of the same pallet so resends cannot become duplicate records, and alerts "
              "before operations notices."],
        ["D", "Close distractor. Fixes the route but blindly resends every stuck message, which would post T1008 and "
              "T1010 as duplicate serial records. With no alerting, failures are retried silently."],
    ], widths=[0.7, 6.5], font=9)
    bullet(doc, "Alerts on retry exhaustion or parked-queue age (for example, more than 5 minutes); turning "
                "dedupe_check on; a synthetic end-to-end test transaction every few minutes; a post-change TLS "
                "handshake and truststore validation step in the proxy patch procedure; certificate-chain monitoring. "
                "Any one is accepted.", bold_lead="Control improvement (any one): ")

    heading(doc, "B. Quick Verification \u2013 Common Wrong Flags")
    para(doc, "Use this table to give targeted feedback when a team submits an incorrect flag.", size=9.5)
    table(doc, ["Submitted Element", "Likely Misconception"], [
        ["HUB", "Blamed the component that reported the 504s rather than the component that produced them"],
        ["TRACECLOUD", "Read the TLS error as a vendor fault; missed the status page and the successful direct hub test"],
        ["FIREWALL", "Followed CHG-4460; missed that all TCP tests succeed and VLAN 30 only affects scanners"],
        ["0906", "Used the hub's first error instead of the proxy's own log"],
        ["0852 / 0855", "Used the change or restart time; missed that pooled connections kept working until max_age"],
        ["8", "Counted every row marked Missing on the dashboard"],
        ["7", "Counted every parked hub message, including both duplicates"],
        ["6", "Removed only one duplicate, or included T1005"],
        ["D", "Fixed the route but did not consider duplicate replay"],
    ], widths=[1.6, 5.6], font=9)

    heading(doc, "C. 45-Minute Calibration (team of 3\u20134)")
    table(doc, ["Time", "Expected Activity", "Author Check"], [
        ["0\u20137 min", "Read scenario, ticket, system flow; assign roles", "Team understands the business impact and splits the evidence"],
        ["7\u201322 min", "Parallel analysis by role", "Reconciliation analyst matches dedupe keys; log analyst finds proxy TLS failure; infrastructure analyst rules out firewall and vendor"],
        ["22\u201330 min", "Regroup and compare findings", "Team reconciles hub 504s with proxy log and works out the max_age onset"],
        ["30\u201338 min", "Evaluate architecture options", "Team rejects D because of duplicate replay and separates immediate fix from future design"],
        ["38\u201343 min", "Complete answer sheet", "Reasoning supports every flag element"],
        ["43\u201345 min", "Build and submit flag", "Flag is exact and verifiable in seconds"],
    ], widths=[0.9, 2.4, 3.9], font=9)

    heading(doc, "D. Revision Log")
    para(doc, "Version 1 \u2192 Version 2 (author self-review against the rubric, before playtesting):", size=10).runs[0].bold = True
    table(doc, ["Issue in Version 1", "Change in Version 2"], [
        ["Formatting rule gave away the answer (\u201cuse LEGACYPROXY if\u2026\u201d)", "Replaced with a neutral table of seven component codes"],
        ["Notes column labelled rows \u201cUnique receipt\u201d and \u201cManual resend of T1004\u201d", "Removed notes; duplicates must be found by matching dedupe keys across Evidence A, B, and G"],
        ["Proxy log said the hub session was accepted, but the connectivity test said hub-to-proxy timed out", "Made all evidence consistent: hub-to-proxy TCP works, proxy-to-cloud TLS fails, hub returns 504"],
        ["Option C was clearly the longest and most detailed choice", "Rewrote options to similar length; added a close distractor (D) that fails only on duplicate replay"],
        ["Could be solved by one person in roughly 10 minutes", "Added four roles, scanner/ops log, change calendar, onset-time element, a non-integration rejection (T1005), and a second duplicate (T1010)"],
        ["Flag had three elements (one was given away)", "Flag has four independent elements; added a common-wrong-flag table for fast feedback"],
        ["Single hint", "Two timed hints that point to method, not answers"],
    ], widths=[3.4, 3.8], font=9)
    para(doc, "Playtest feedback \u2192 Version 3 (complete after playtesting):", size=10).runs[0].bold = True
    table(doc, ["Feedback Received", "Change Made", "Status"], [["", "", ""], ["", "", ""], ["", "", ""]],
          widths=[3.0, 3.2, 1.0], font=9)

    heading(doc, "E. Playtest Feedback Form")
    para(doc, "Give this section to the playtest team after they submit a flag.", size=9.5)
    table(doc, ["Field", "Response"], [
        ["Playtest team / members", ""], ["Date", ""], ["Start time / finish time", ""],
        ["Flag submitted", ""], ["Correct on first attempt? (Y/N)", ""], ["Hints used (0 / 1 / 2)", ""],
    ], widths=[2.4, 4.8], font=9.5)
    for q in [
        "1. How many minutes did the challenge take your team to complete?",
        "2. Were the instructions clear enough to begin without additional explanation? Why or why not?",
        "3. Did the role split work? Did every team member have meaningful work?",
        "4. Which evidence item was most useful in identifying the root cause?",
        "5. Was any clue confusing, unnecessary, contradictory, or too obvious?",
        "6. Could any part of the flag be guessed without completing the analysis?",
        "7. Did the challenge make you distinguish between the immediate fix and the long-term preventive design?",
        "8. Did the amount of evidence feel appropriate for a 45-minute challenge?",
        "9. What did you learn about real-time data integration or troubleshooting?",
        "10. What is one change that would make the challenge clearer or more educational?",
    ]:
        answer_lines(doc, q, 2)

    heading(doc, "F. Guide for Giving Feedback on Another Team\u2019s Challenge")
    for txt in [
        "Start with one specific strength: realism, clarity, evidence quality, collaboration, or verification design.",
        "Comment on whether the challenge required real analysis rather than guessing or recalling definitions.",
        "State whether the actual duration was close to 45 minutes, and where your team spent the most time.",
        "Note whether the work naturally split across team members or could be done by one person.",
        "Identify one area that could be clearer, less complex, or more challenging.",
        "Check whether the final answer proves the team actually solved the challenge.",
        "Suggest one practical improvement, such as adding an evidence source, adjusting a hint, or removing a giveaway.",
    ]:
        bullet(doc, txt)

    heading(doc, "G. Debrief Synopsis (200\u2013300 words)")
    para(doc, "Draft. Replace the bracketed sentences with what actually happened in your playtest before submitting.",
         italic=True, color=RED, size=9.5)
    for lead, txt in DEBRIEF:
        para(doc, txt, bold_lead=lead, size=10.5)

    heading(doc, "H. Final Submission Checklist")
    for txt in [
        "Challenge still takes about 45 minutes after revisions.",
        "Instructions can be followed without verbal explanation.",
        "No answer is revealed in the participant packet.",
        "The flag requires findings from at least three different evidence sources.",
        "Immediate recovery and long-term prevention are both addressed.",
        "Playtest feedback is recorded in Section D and incorporated where useful.",
        "Debrief reflects the actual playtest and is 200\u2013300 words.",
        "Team leader submits only the final revised version (packet + author materials).",
    ]:
        bullet(doc, "\u2610  " + txt, style="Normal")

    path = OUT_DIR + r"\Assignment_6.6_Author_Key_Playtest_Debrief_MS_v2.docx"
    doc.save(path)
    return path


DEBRIEF = [
    ("Experience. ",
     "I chose Real-Time Data Integration and Process Optimization because integration failures are among the most "
     "common problems in day-to-day IT support, and they are rarely solved by looking at one system. Turning a "
     "real troubleshooting pattern into a CTF was harder than expected. Every log, table, and timestamp had to agree "
     "with each other and still lead to exactly one verifiable flag."),
    ("Lessons learned. ",
     "The biggest lesson was that the component reporting an error is often not the component that caused it. The "
     "integration hub logged the timeouts, but the legacy proxy had failed minutes earlier after a routine patch "
     "rebuilt its certificate truststore. I also saw how easily duplicate data enters a system during an outage, when "
     "well-meaning staff resend or rescan, and why idempotency must be in place before anything is replayed. On the "
     "design side, my first draft gave the answer away through a formatting note and a notes column, and two pieces "
     "of evidence contradicted each other. Writing a table of common wrong flags forced me to confirm that each clue "
     "supported only one answer."),
    ("Playtest feedback and revision. ",
     "[Playtest team] completed the challenge in [__] minutes and submitted [flag]. Their main feedback was "
     "[feedback]. Based on that, I [revision made]."),
    ("Final thought. ",
     "In future BSA and integration work, I will trace data end to end, check recent changes first, and confirm "
     "that a recovery plan cannot create duplicates before replaying anything."),
]


if __name__ == "__main__":
    print(build_packet())
    print(build_author())
    words = sum(len(t.split()) for _, t in DEBRIEF)
    print("debrief words (incl. placeholders):", words)
