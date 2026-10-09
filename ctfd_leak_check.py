import docx

from ctfd_session import BASE, login

LEAKS = ["LEGACYPROXY_0905", "Answer Key", "Author Materials", "Correct flag", "Common Wrong Flags",
         "T1003, T1004, T1006, T1007, T1009", "Option C.", "Close distractor"]

s = login()
cid = next(c["id"] for c in s.get(BASE + "/api/v1/challenges").json()["data"] if c["name"] == "The Missing Trace")
c = s.get(f"{BASE}/api/v1/challenges/{cid}").json()["data"]
player_text = c["description"] + " ".join(h["content"] for h in s.get(f"{BASE}/api/v1/challenges/{cid}/hints").json()["data"])
print("site files:", c["files"])
print("site leaks:", [x for x in LEAKS if x in player_text])

d = docx.Document("Assignment_6.6_CTF_Challenge_Packet_MS_v2.docx")
packet = "\n".join(p.text for p in d.paragraphs) + "\n".join(cell.text for t in d.tables for r in t.rows for cell in r.cells)
print("packet leaks:", [x for x in LEAKS if x in packet])
print("user mode:", s.get(BASE + "/api/v1/configs/user_mode").json()["data"]["value"])
